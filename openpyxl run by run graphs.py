##import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Series, Reference
import plotly.plotly as py
import os
##py.sign_in('barryguglielmo','WNMyBgUv3C0tHQzvfN33')
os.chdir('C:/Users/bag019/Desktop/data graphs')
wb = load_workbook('Database Fo Real.xlsx',data_only=True)

##raw data
tabs = ['WP AI', 'WP HX', 'PPT E', 'TG', 'HDL-C', 'TC', 'AI (P1)', 'C3 (P1)', 'HP (P1)', 'E (P1)', 'J #1 (P1)', 'C3 PPT (P1)', 'J  # 2 (P1)', 'J PPT (P1)', 'AI(C3+) (P3)', 'AI(HP+) (P3)', 'AI(E+) (P3)', 'AI(J+) #1 (P3)', 'E(AI+) (P3)', 'E(C3+) PPT (P3)', 'E(J+) #2 (P3)']
def rawbar (in_sheet):
   "What raw data sheet to chart"
   counter = 3
   r = 1
   loci = 3
   ws_from = wb.get_sheet_by_name(in_sheet)
   while counter < 1803:
      chart1 = BarChart()
      chart1.type = "col"
      chart1.style = 6
      chart1.title = "Raw Data Run " + str(r)
      chart1.y_axis.title = 'Absorbance'
      chart1.x_axis.title = 'Sample I.D.'

      data = Reference(ws_from, min_col=11, min_row=counter, max_row= counter + 35)

      chart1.add_data(data)
      chart1.shape = 4

      my_graph = wb.get_sheet_by_name(in_sheet)
      my_graph.add_chart(chart1, 'AQ'+ str(loci))
      counter = counter + 36
      r = r+ 1
      loci = loci + 36
def contbar (in_sheet):
   "What control data sheet to chart"
   counter = 3
   r = 1
   loci = 3
   ws_from = wb.get_sheet_by_name(in_sheet)
   while counter < 1803:
      chart1 = BarChart()
      chart1.type = "col"
      chart1.style = 6
      chart1.title = "Control Data Run " + str(r)
      chart1.y_axis.title = 'Absorbance'
      chart1.x_axis.title = 'Sample I.D.'

      data = Reference(ws_from, min_col=39, min_row=counter, max_row= counter + 35)

      chart1.add_data(data)
      chart1.shape = 4

      my_graph = wb.get_sheet_by_name(in_sheet)
      my_graph.add_chart(chart1, 'AZ' + str(loci))
      counter = counter + 36
      r = r+ 1
      loci = loci + 36

for tab in tabs:
   rawbar(tab)
   contbar(tab)

wb.save('Database For Run Analysis.xlsx') 
## need to know how to play with x axis
